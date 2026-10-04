"""Generate 5 synthetic portfolio companies x 6 months of deliberately messy reports.

All data is SYNTHETIC. Outputs:
  data/raw/<company>/<period>.<ext>   the messy monthly files
  data/truth.csv                      ground truth in canonical USD (for scoring the pipeline)
  data/planted_errors.csv             the errors we injected (for measuring DQ recall)
"""
import csv
import random
from pathlib import Path

import openpyxl
import yaml
from openpyxl.styles import Alignment, Font
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Table, TableStyle

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "data" / "raw"
MONTHS = [f"2026-{m:02d}" for m in range(1, 7)]
MULT = {"thousands": 1e3, "millions": 1e6, "lakhs": 1e5, "crore": 1e7}
FX = yaml.safe_load((ROOT / "config" / "fx.yaml").read_text())["rates_to_usd"]
MON = ["Jan", "Feb", "Mar", "Apr", "May", "Jun"]

# label per canonical KPI, per company. `drift_label` = (month_index, kpi, new_label)
COMPANIES = {
    "nimbus_cloud": dict(
        fmt="xlsx", cur="USD", unit="thousands", base_rev=420_000, saas=True,
        labels=dict(revenue="Net Sales", cogs="Cost of Sales", gross_profit="Gross Profit",
                    ebitda="EBITDA", cash="Cash & Equivalents", headcount="FTEs",
                    arr="ARR", churn_pct="Monthly Churn %"),
        reorder_months=[2, 4]),
    "gulf_logistics": dict(
        fmt="csv", cur="AED", unit="millions", base_rev=1_900_000, saas=False,
        labels=dict(revenue="Turnover", cogs="Cost of Goods Sold", gross_profit="Gross Profit",
                    ebitda="Operating Profit", cash="Cash Balance", headcount="Headcount"),
        reorder_months=[3]),
    "bharat_foods": dict(
        fmt="pdf", cur="INR", unit="lakhs", base_rev=650_000, saas=False,
        labels=dict(revenue="Revenue from Operations", gross_margin_pct="Gross Margin (%)",
                    ebitda="EBITDA", cash="Cash and Bank", headcount="Employees"),
        reorder_months=[1, 5]),
    "desert_solar": dict(
        fmt="xlsx", cur="AED", unit="thousands", base_rev=780_000, saas=False, fiscal=True,
        labels=dict(revenue="Total Revenue", cogs="Direct Costs", gross_profit="Gross Profit",
                    ebitda="Operating Profit (EBITDA)", cash="Cash", headcount="No. of Employees"),
        reorder_months=[2]),
    "pune_health": dict(
        fmt="csv", cur="INR", unit="crore", base_rev=310_000, saas=True,
        labels=dict(revenue="Revenue", gross_margin_pct="Gross Margin %", ebitda="EBITDA",
                    cash="Cash", headcount="Headcount", arr="Annual Recurring Revenue",
                    churn_pct="Customer Churn %"),
        reorder_months=[4], drift_label=(2, "churn_pct", "Logo Churn")),
}

# (company, month_index, kpi, how, note) -- injected into REPORTED values only
PLANTED = [
    ("gulf_logistics", 3, "revenue", "x10", "10x typo on revenue"),
    ("nimbus_cloud", 2, "gross_profit", "tie_break", "gross profit does not equal revenue - COGS"),
    ("bharat_foods", 4, "ebitda", "gt_revenue", "EBITDA reported above revenue"),
    ("pune_health", 5, "churn_pct", "set:45.0", "implausible churn of 45%"),
    ("desert_solar", 1, "cash", "null", "cash left blank"),
]


def period_label(co: str, cfg: dict, i: int) -> str:
    """Each company writes the period its own way."""
    if cfg.get("fiscal"):  # FY ends March: Jan-Mar 2026 = FY26 M10-M12, then Apr 2026 starts FY27 M1
        n = i + 10
        return f"FY26 M{n}" if n <= 12 else f"FY27 M{n - 12}"
    if co == "nimbus_cloud":
        return f"{MON[i]} 2026"
    return [f"{MON[i]}-2026", f"2026-{i + 1:02d}", f"{i + 1:02d}/2026"][i % 3]


def truth_series(rng: random.Random, cfg: dict) -> list[dict]:
    rows, cash = [], cfg["base_rev"] * rng.uniform(2.5, 4)
    hc0 = rng.uniform(60, 90)  # headcount drifts slowly, like a real company
    for i in range(len(MONTHS)):
        rev = cfg["base_rev"] * (1.025 ** i) * rng.uniform(0.97, 1.03)
        gm = rng.uniform(0.55, 0.75) if cfg["saas"] else rng.uniform(0.22, 0.38)
        ebitda_m = rng.uniform(0.04, 0.16) if not cfg["saas"] else rng.uniform(-0.05, 0.12)
        cash += rev * ebitda_m * 0.7
        rows.append(dict(
            revenue=rev, cogs=rev * (1 - gm), gross_profit=rev * gm, gross_margin_pct=gm * 100,
            ebitda=rev * ebitda_m, cash=cash, headcount=int(hc0 + i * 1.5 + rng.uniform(-1.5, 1.5)),
            arr=rev * 12 * 1.04, churn_pct=rng.uniform(0.9, 2.4)))
    return rows


def to_native(kpi: str, usd_val: float, cfg: dict) -> float:
    if kpi in ("revenue", "cogs", "gross_profit", "ebitda", "cash", "arr"):
        return usd_val / FX[cfg["cur"]] / MULT[cfg["unit"]]
    return usd_val  # headcount, percentages


def reported_row(co, cfg, i, truth):
    kpis = list(cfg["labels"])
    # gross_profit is derived for tie checks; keep native values rounded as a human would
    vals = {k: round(to_native(k, truth[k], cfg), 2 if k != "headcount" else 0) for k in kpis}
    for (pco, pi, kpi, how, _) in PLANTED:
        if pco == co and pi == i and kpi in vals:
            if how == "x10":
                vals[kpi] = round(vals[kpi] * 10, 2)
            elif how == "tie_break":
                vals[kpi] = round(vals[kpi] * 1.12, 2)
            elif how == "gt_revenue":
                vals[kpi] = round(vals["revenue"] * 1.4, 2)
            elif how.startswith("set:"):
                vals[kpi] = float(how.split(":")[1])
            elif how == "null":
                vals[kpi] = None
    return vals


def header_order(co, cfg, i):
    """Column order drifts in the configured months."""
    kpis = list(cfg["labels"])
    if i in cfg["reorder_months"]:
        r = random.Random(f"{co}-{i}")
        r.shuffle(kpis)
    return kpis


def label_for(cfg, kpi, i):
    d = cfg.get("drift_label")
    if d and kpi == d[1] and i >= d[0]:
        return d[2]
    return cfg["labels"][kpi]


def note(cfg):
    return f"All figures in {cfg['cur']} {cfg['unit']} unless stated otherwise"


def write_csv(path, co, cfg, i, kpis, vals):
    with path.open("w", newline="") as f:
        w = csv.writer(f)
        f.write(f"# {co.replace('_', ' ').title()} - monthly management report\n")
        f.write(f"# {note(cfg)}\n")
        w.writerow(["Period"] + [label_for(cfg, k, i) for k in kpis])
        w.writerow([period_label(co, cfg, i)] + ["" if vals[k] is None else vals[k] for k in kpis])


def write_xlsx(path, co, cfg, i, kpis, vals):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Report"
    ncol = len(kpis) + 1
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
    ws["A1"] = f"{co.replace('_', ' ').title()} | Monthly KPI Pack"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncol)
    ws["A2"] = note(cfg)
    ws.merge_cells(start_row=4, start_column=2, end_row=4, end_column=ncol)
    ws["B4"] = "Key metrics"  # merged group header above the real header
    ws["B4"].alignment = Alignment(horizontal="center")
    ws.cell(5, 1, "Period")
    ws.cell(6, 1, period_label(co, cfg, i))
    for c, k in enumerate(kpis, start=2):
        ws.cell(5, c, label_for(cfg, k, i)).font = Font(bold=True)
        ws.cell(6, c, vals[k])
    wb.save(path)


def write_pdf(path, co, cfg, i, kpis, vals):
    styles = getSampleStyleSheet()
    doc = SimpleDocTemplate(str(path), pagesize=landscape(A4))
    head = [Paragraph(f"<b>{co.replace('_', ' ').title()}</b> - Monthly Board Pack", styles["Title"]),
            Paragraph(note(cfg), styles["Normal"])]
    fmt = lambda k, v: "" if v is None else (f"{v:,.0f}" if k == "headcount" else f"{v:,.2f}")
    data = [["Period"] + [label_for(cfg, k, i) for k in kpis],
            [period_label(co, cfg, i)] + [fmt(k, vals[k]) for k in kpis]]
    t = Table(data)
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.grey), ("FONTSIZE", (0, 0), (-1, -1), 8)]))
    doc.build(head + [t])


WRITERS = {"csv": write_csv, "xlsx": write_xlsx, "pdf": write_pdf}


def main(seed: int = 42):
    rng = random.Random(seed)
    truth_rows, all_series = [], {}
    for co, cfg in COMPANIES.items():
        (RAW / co).mkdir(parents=True, exist_ok=True)
        series = truth_series(rng, cfg)
        all_series[co] = series
        for i, month in enumerate(MONTHS):
            vals = reported_row(co, cfg, i, series[i])
            kpis = header_order(co, cfg, i)
            WRITERS[cfg["fmt"]](RAW / co / f"{month}.{cfg['fmt']}", co, cfg, i, kpis, vals)
            for k in cfg["labels"]:
                truth_rows.append(dict(company=co, month=month, kpi=k,
                                       usd_value=round(series[i][k], 4)))
    out = ROOT / "data"
    with (out / "truth.csv").open("w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=["company", "month", "kpi", "usd_value"])
        w.writeheader(); w.writerows(truth_rows)
    with (out / "planted_errors.csv").open("w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["company", "month", "kpi", "kind", "note"])
        for co, i, kpi, how, n in PLANTED:
            w.writerow([co, MONTHS[i], kpi, how, n])
    # What an analyst would get back after querying the company: corrected native-unit figures.
    resub = [dict(company=co, month=MONTHS[i], kpi=kpi,
                  corrected_value=round(to_native(kpi, all_series[co][i][kpi], COMPANIES[co]), 0 if kpi == "headcount" else 2),
                  reason="Company confirmed and resubmitted the figure (synthetic scenario)",
                  approved_by="analyst (synthetic)")
             for co, i, kpi, how, n in PLANTED]
    (out / "resubmissions.yaml").write_text(yaml.safe_dump(resub, sort_keys=False))
    print(f"Generated {sum(len(list((RAW / c).glob('*'))) for c in COMPANIES)} files in {RAW}")


if __name__ == "__main__":
    main()
